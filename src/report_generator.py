"""
src/report_generator.py
=======================
Modul Pembangkit Laporan Audit Steganalisis Forensik Format Microsoft Word (.docx)
Menggunakan font Times New Roman, tata letak tabel rapi berstandar laporan akademik.
"""

import io
from typing import Dict, Any
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn


def set_cell_background(cell, color_hex: str):
    """Mengatur warna latar belakang sel tabel (shading)."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Mengatur padding sel tabel."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def create_forensic_docx_report(report_data: Dict[str, Any]) -> io.BytesIO:
    """
    Menghasilkan berkas Word .docx berisi Laporan Audit Steganalisis Forensik
    dengan font Times New Roman dan tata letak profesional.
    """
    doc = docx.Document()

    # Atur margin halaman standar 1 inci (2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Gaya dasar font Times New Roman
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    # ==========================================
    # 1. KOP / JUDUL LAPORAN
    # ==========================================
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("LAPORAN AUDIT STEGANALISIS FORENSIK CITRA")
    run_title.font.name = "Times New Roman"
    run_title.font.size = Pt(15)
    run_title.font.bold = True

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run("HideBit Forensic Telemetry System — Evaluasi Keamanan Informasi Citra")
    run_sub.font.name = "Times New Roman"
    run_sub.font.size = Pt(10.5)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    # ==========================================
    # 2. METADATA PEMERIKSAAN
    # ==========================================
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(10)
    h1.paragraph_format.space_after = Pt(4)
    r_h1 = h1.add_run("I. IDENTITAS & METADATA PEMERIKSAAN")
    r_h1.font.name = "Times New Roman"
    r_h1.font.size = Pt(12)
    r_h1.font.bold = True

    meta_table = doc.add_table(rows=0, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    metadata_items = [
        ("Tanggal & Waktu Pemeriksaan", str(report_data.get("timestamp", "-"))),
        ("Metode Pengujian", str(report_data.get("mode", "Blind Steganalysis (Single Image)"))),
        ("Nama Berkas Citra Uji", str(report_data.get("fileName", "-"))),
        ("Ukuran Berkas Uji", str(report_data.get("fileSize", "-"))),
    ]

    has_comp = bool(report_data.get("comparison") and report_data["comparison"].get("has_reference"))
    if has_comp:
        metadata_items.append(("Citra Referensi Cover", str(report_data.get("refName", "Cover Asli"))))

    col_widths = [Inches(2.5), Inches(4.0)]
    for label, val in metadata_items:
        row = meta_table.add_row()
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width, cell_val.width = col_widths[0], col_widths[1]

        p_lbl = cell_lbl.paragraphs[0]
        p_lbl.paragraph_format.space_after = Pt(2)
        r = p_lbl.add_run(label)
        r.font.name = "Times New Roman"
        r.font.size = Pt(10.5)
        r.font.bold = True

        p_val = cell_val.paragraphs[0]
        p_val.paragraph_format.space_after = Pt(2)
        r2 = p_val.add_run(f": {val}")
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(10.5)

        set_cell_margins(cell_lbl, top=40, bottom=40, left=80, right=80)
        set_cell_margins(cell_val, top=40, bottom=40, left=80, right=80)

    # ==========================================
    # 3. HASIL DIAGNOSIS FORENSIK
    # ==========================================
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(4)
    r_h2 = h2.add_run("II. HASIL DIAGNOSIS & STATUS FORENSIK")
    r_h2.font.name = "Times New Roman"
    r_h2.font.size = Pt(12)
    r_h2.font.bold = True

    score = float(report_data.get("score", 0.0))
    verdict = report_data.get("verdict", {})
    badge_label = verdict.get("badge_label", "MENUNGGU HASIL")
    title_diag = verdict.get("title", "-")
    narrative = verdict.get("narrative", "-")
    recommendation = verdict.get("recommendation", "-")

    box_table = doc.add_table(rows=1, cols=1)
    box_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    box_cell = box_table.rows[0].cells[0]
    box_cell.width = Inches(6.5)
    set_cell_background(box_cell, "F4F6F9" if score < 40 else "FFF3CD" if score < 65 else "F8D7DA")
    set_cell_margins(box_cell, top=120, bottom=120, left=180, right=180)

    p_box = box_cell.paragraphs[0]
    p_box.paragraph_format.space_after = Pt(4)
    r_sc = p_box.add_run(f"Skor Kecurigaan Heuristik: {score:.2f} / 100 — Status: {badge_label}\n")
    r_sc.font.name = "Times New Roman"
    r_sc.font.size = Pt(11)
    r_sc.font.bold = True
    if score >= 65:
        r_sc.font.color.rgb = RGBColor(0xA0, 0x00, 0x00)
    elif score >= 40:
        r_sc.font.color.rgb = RGBColor(0x85, 0x64, 0x04)
    else:
        r_sc.font.color.rgb = RGBColor(0x15, 0x57, 0x24)

    r_ti = p_box.add_run(f"Diagnosis: {title_diag}\n\n")
    r_ti.font.name = "Times New Roman"
    r_ti.font.size = Pt(10.5)
    r_ti.font.bold = True

    r_na = p_box.add_run(f"Temuan Analisis:\n{narrative}\n\n")
    r_na.font.name = "Times New Roman"
    r_na.font.size = Pt(10.5)

    r_re = p_box.add_run(f"Rekomendasi Tindakan:\n{recommendation}")
    r_re.font.name = "Times New Roman"
    r_re.font.size = Pt(10.5)
    r_re.font.italic = True

    # ==========================================
    # 4. PROFIL STATISTIK CHI-SQUARE PoV (PAIRS OF VALUES)
    # ==========================================
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(4)
    r_h3 = h3.add_run("III. PROFIL STATISTIK CHI-SQUARE PAIRS OF VALUES (PoV)")
    r_h3.font.name = "Times New Roman"
    r_h3.font.size = Pt(12)
    r_h3.font.bold = True

    p_pov_desc = doc.add_paragraph()
    p_pov_desc.paragraph_format.space_after = Pt(6)
    r_pvd = p_pov_desc.add_run(
        "Berdasarkan metode steganografi spasial klasik (Westfeld & Pfitzmann, 1999), "
        "uji Pairs of Values mengevaluasi kesetimbangan frekuensi pasangan nilai piksel (2k, 2k+1) "
        "yang cenderung seimbang apabila disusupi bit-bit acak terenkripsi."
    )
    r_pvd.font.name = "Times New Roman"
    r_pvd.font.size = Pt(10)
    r_pvd.font.italic = True

    pov_table = doc.add_table(rows=1, cols=6)
    pov_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    pov_headers = ["Kanal", "Chi-Square (χ²)", "Degrees of Freedom", "p-value", "Kesetimbangan PoV", "Status Kanal"]
    col_w = [Inches(1.0), Inches(1.3), Inches(1.1), Inches(1.0), Inches(1.1), Inches(1.0)]

    hdr_row = pov_table.rows[0]
    for idx, text in enumerate(pov_headers):
        cell = hdr_row.cells[idx]
        cell.width = col_w[idx]
        set_cell_background(cell, "2D3748")
        set_cell_margins(cell, top=80, bottom=80, left=60, right=60)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    channels = report_data.get("channels", {})
    chan_order = [("Red", channels.get("red", {})), ("Green", channels.get("green", {})), ("Blue", channels.get("blue", {}))]

    for ch_name, ch_data in chan_order:
        row = pov_table.add_row()
        chi2_val = ch_data.get("chi_square")
        chi2_str = f"{chi2_val:,.2f}" if chi2_val is not None else "--"
        dof_str = str(ch_data.get("degrees_of_freedom", "--"))
        
        pval_val = ch_data.get("p_value")
        if pval_val is not None:
            pval_str = "< 0.0001" if pval_val < 0.0001 else f"{pval_val:.4f}"
        else:
            pval_str = "--"
            
        bal_str = f"{ch_data.get('balance_percent', '--')}%"
        status_str = str(ch_data.get("status", "Normal"))

        row_vals = [ch_name, chi2_str, dof_str, pval_str, bal_str, status_str]
        for c_idx, val in enumerate(row_vals):
            cell = row.cells[c_idx]
            cell.width = col_w[c_idx]
            set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx > 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.name = "Times New Roman"
            r.font.size = Pt(9.5)
            if c_idx == 0 or c_idx == 5:
                r.font.bold = True

    # ==========================================
    # 5. EVALUASI KOMPARATIF (JIKA ADA REFERENSI COVER)
    # ==========================================
    if has_comp:
        h4 = doc.add_paragraph()
        h4.paragraph_format.space_before = Pt(14)
        h4.paragraph_format.space_after = Pt(4)
        r_h4 = h4.add_run("IV. HASIL EVALUASI KOMPARATIF (CITRA COVER VS STEGO)")
        r_h4.font.name = "Times New Roman"
        r_h4.font.size = Pt(12)
        r_h4.font.bold = True

        comp_metrics = report_data["comparison"].get("metrics", {})
        comp_table = doc.add_table(rows=1, cols=3)
        comp_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        c_headers = ["Parameter Metrik", "Nilai Kuantitatif", "Keterangan Evaluasi"]
        c_widths = [Inches(2.2), Inches(1.8), Inches(2.5)]

        c_hdr_row = comp_table.rows[0]
        for idx, text in enumerate(c_headers):
            cell = c_hdr_row.cells[idx]
            cell.width = c_widths[idx]
            set_cell_background(cell, "4A5568")
            set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.name = "Times New Roman"
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        comp_rows = [
            ("PSNR (Peak Signal-to-Noise Ratio)", f"{comp_metrics.get('psnr_db', '--')} dB", "Ambang batas minimal kelayakan: ≥ 30 dB"),
            ("MSE (Mean Squared Error)", f"{comp_metrics.get('mse', '--')}", "Nilai mendekati 0 menunjukkan distorsi sangat rendah"),
            ("Piksel Dimodifikasi", f"{comp_metrics.get('changed_pixels', 0):,} piksel ({comp_metrics.get('pixel_change_percent', 0)}%)", "Sebaran merata akibat permutasi PRNG"),
            ("Piksel Identik (Utuh)", f"{comp_metrics.get('unchanged_pixels', 0):,} piksel ({comp_metrics.get('unchanged_percent', 0)}%)", "Piksel yang mempertahankan nilai asli"),
            ("Maksimum Delta Nilai Piksel", f"± {comp_metrics.get('max_delta', 1)}", "Konsisten dengan karakteristik substitusi 1-bit LSB"),
        ]

        for p_lbl, p_val, p_ket in comp_rows:
            row = comp_table.add_row()
            cells = row.cells
            cells[0].width, cells[1].width, cells[2].width = c_widths[0], c_widths[1], c_widths[2]
            
            p0 = cells[0].paragraphs[0]
            p0.paragraph_format.space_after = Pt(0)
            r0 = p0.add_run(p_lbl)
            r0.font.name = "Times New Roman"
            r0.font.size = Pt(9.5)
            r0.font.bold = True

            p1 = cells[1].paragraphs[0]
            p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p1.paragraph_format.space_after = Pt(0)
            r1 = p1.add_run(p_val)
            r1.font.name = "Times New Roman"
            r1.font.size = Pt(9.5)

            p2 = cells[2].paragraphs[0]
            p2.paragraph_format.space_after = Pt(0)
            r2 = p2.add_run(p_ket)
            r2.font.name = "Times New Roman"
            r2.font.size = Pt(9.5)

            for cell in cells:
                set_cell_margins(cell, top=60, bottom=60, left=80, right=80)

    # ==========================================
    # 6. CATATAN PENUTUP & AUTENTIKASI
    # ==========================================
    doc.add_paragraph().paragraph_format.space_before = Pt(14)
    p_close = doc.add_paragraph()
    p_close.paragraph_format.space_after = Pt(2)
    r_cl = p_close.add_run(
        "Laporan audit steganografi dan forensik citra ini diterbitkan secara otomatis oleh engine HideBit "
        "sebagai luaran verifikasi integritas data dan analisis keamanan informasi."
    )
    r_cl.font.name = "Times New Roman"
    r_cl.font.size = Pt(9.5)
    r_cl.font.italic = True
    r_cl.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    output_stream = io.BytesIO()
    doc.save(output_stream)
    output_stream.seek(0)
    return output_stream
